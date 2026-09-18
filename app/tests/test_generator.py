import hashlib
import io
import json
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apx_helper import contracts
from apx_helper.generator import render, render_library, render_transaction, render_transactions
from apx_helper.settings import (
    LIBRARY_TEMPLATE_GLOB,
    TRANSACTION_TEMPLATE_GLOB,
    find_template,
    templates_dir,
)

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def _load(name: str) -> dict:
    return json.loads((EXAMPLES / name).read_text(encoding="utf-8"))


@pytest.fixture
def transaction() -> dict:
    return _load("ejemplo-transaccion.json")


@pytest.fixture
def library() -> dict:
    return _load("ejemplo-libreria.json")


def test_detecta_tipo(transaction, library):
    assert contracts.detect_kind(transaction) == contracts.TRANSACTION
    assert contracts.detect_kind(library) == contracts.LIBRARY


def test_rechaza_estructura_alterada(transaction):
    invalid = deepcopy(transaction)
    invalid["TECHNICAL DATA"].pop("Country")
    invalid["EXTRA"] = "x"
    with pytest.raises(contracts.ContractError) as exc:
        contracts.validate(invalid)
    assert "Country" in str(exc.value)
    assert "EXTRA" in str(exc.value)


def test_transaccion_escribe_campos_y_tablas(transaction):
    result = render_transaction(transaction)
    ws = result.workbook.worksheets[0]

    assert ws.title == "ADVST501-01-AR"
    assert result.filename == "ADVST501-01-AR - APX Transactions Global Sheet v1.1.xlsx"
    assert ws["B7"].value.startswith("Recupera")
    assert ws["C17"].value == "AR"
    assert ws["F17"].value == "01"
    assert ws["C21"].value == "No"
    assert ws["F21"].value == "Yes"
    assert ws["C29"].value == "ADVST501-01-AR"
    assert ws["C34"].value == "ADVSR500"
    assert ws["D35"].value == "executeGetDeviceCatalog"
    assert ws["D40"].value == "06 - ERROR NO ROLLBACK"
    assert ws["B52"].value == "Input"
    assert ws["C52"].value == "customerId"
    assert ws["E53"].value == "date"
    assert ws["B60"].value == "Output"
    assert ws["E61"].value == "long"


def test_transaccion_marca_valores_no_validos(transaction):
    transaction["TECHNICAL DATA"]["Country"] = "[REVISAR: BR]"
    result = render_transaction(transaction)
    ws = result.workbook.worksheets[0]

    assert ws["C17"].fill.start_color.rgb == "FFFFE699"
    assert any("Country" in warning for warning in result.warnings)


def test_transaccion_inserta_filas_si_sobran_elementos(transaction):
    extra = [
        {"Name of APX field": f"field{i}", "Mandatory?": "Yes", "APX data type": "string", "Description": f"campo {i}"}
        for i in range(8)
    ]
    transaction["PARAMETERS"]["Input Parameters"] = extra
    result = render_transaction(transaction)
    ws = result.workbook.worksheets[0]

    assert ws["C52"].value == "field0"
    assert ws["C59"].value == "field7"
    # Los parámetros de salida se desplazan 3 filas y conservan su encabezado.
    assert ws["C62"].value == "Name of APX field"
    assert ws["C63"].value == "events"
    assert "C63" in {str(rng) for rng in ws.merged_cells.ranges} or ws["F63"].value is not None


def test_libreria_genera_una_hoja_por_metodo(library):
    result = render_library(library)
    wb = result.workbook

    assert wb.sheetnames[:3] == ["ADVSR500", "executeSaveBiometricEvents", "executeGetBiometricEvents"]
    assert "executeMethod" not in wb.sheetnames
    assert result.filename == "ADVSR500 - APX Libraries Global Sheet V1.1.xlsx"

    main = wb["ADVSR500"]
    assert main["C15"].value == "Public"
    assert main["F15"].value == "On-line"
    assert main["B23"].value == "executeSaveBiometricEvents"
    assert main["C23"].value == "Persiste una lista de eventos biométricos."
    assert main["B24"].value == "executeGetBiometricEvents"

    method = wb["executeGetBiometricEvents"]
    assert method["B2"].value == "executeGetBiometricEvents"
    assert method["C16"].value == "MongoDB"
    assert method["C20"].value == "ADVS00060020"
    assert method["C24"].value == "Consulta de eventos biométricos"
    assert method["E24"].value == "ADVS-EVT-QUERY"
    assert method["B31"].value == "Input"
    assert method["C31"].value == "customerId"
    assert method["C39"].value == "biometricEventsDTO"


def test_libreria_inserta_filas_en_bloque_de_errores(library):
    result = render_library(library, library_id="ADVSR500")
    method = result.workbook["executeSaveBiometricEvents"]
    # 3 errores en un bloque con capacidad 2 -> se inserta una fila.
    assert method["C20"].value == "ADVS00060010"
    assert method["C22"].value == "ADVS00040012"
    # Todo lo que había debajo baja una fila conservando encabezados y combinaciones.
    assert method["C24"].value == "Functional name"
    assert "C25:D25" in {str(rng) for rng in method.merged_cells.ranges}
    assert method["B29"].value == "PARAMETERS"


def test_render_guarda_archivo(tmp_path, transaction):
    result = render(transaction)
    destination = result.save(tmp_path)
    assert destination.exists() and destination.stat().st_size > 0


def test_varias_transacciones_en_un_solo_libro(transaction):
    segunda = deepcopy(transaction)
    segunda["TECHNICAL DATA"]["Transaction Identifier"] = "ADVST511-01-AR"
    segunda["TECHNICAL DATA"]["Country"] = "MX"
    segunda["PARAMETERS"]["Input Parameters"] = [
        {"Name of APX field": f"campo{i}", "Mandatory?": "Yes", "APX data type": "string", "Description": "x"}
        for i in range(7)
    ]

    result = render_transactions([transaction, segunda])
    wb = result.workbook

    assert wb.sheetnames == [
        "ADVST501-01-AR",
        "ADVST511-01-AR",
        "Parameters (Do not remove)",
        "Change Log",
    ]
    assert result.filename == "ADVST501-01-AR - APX Transactions Global Sheet v1.1.xlsx"
    # Cada hoja se rellena de forma independiente: la primera no hereda filas de la segunda.
    assert wb["ADVST501-01-AR"]["C17"].value == "AR"
    assert wb["ADVST501-01-AR"]["B59"].value == "Output Parameters"
    assert wb["ADVST511-01-AR"]["C17"].value == "MX"
    assert wb["ADVST511-01-AR"]["C58"].value == "campo6"
    assert wb["ADVST511-01-AR"]["B61"].value == "Output Parameters"


def test_render_rechaza_mezclar_tipos(transaction, library):
    with pytest.raises(contracts.ContractError):
        render([transaction, library])


def _sheet_values(ws) -> list[list]:
    return [[cell.value for cell in row] for row in ws.iter_rows()]


@pytest.mark.parametrize(
    "glob, render_call",
    [
        (TRANSACTION_TEMPLATE_GLOB, lambda data: render_transactions([data, data])),
        (LIBRARY_TEMPLATE_GLOB, lambda data: render_library(data, library_id="LIBTEST")),
    ],
)
def test_hojas_fijas_de_la_plantilla_no_mutan(glob, render_call, transaction, library):
    """'Parameters (Do not remove)' y 'Change Log' viajan intactas y el archivo maestro no se toca."""
    template = find_template(glob)
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    original = load_workbook(template)

    data = transaction if glob is TRANSACTION_TEMPLATE_GLOB else library
    generated = load_workbook(io.BytesIO(render_call(data).to_bytes()))

    for name in ("Parameters (Do not remove)", "Change Log"):
        assert _sheet_values(generated[name]) == _sheet_values(original[name])

    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_save_no_escribe_en_la_carpeta_de_plantillas(transaction):
    with pytest.raises(ValueError):
        render_transaction(transaction).save(templates_dir())
