"""Utilidad de inspeccion: vuelca la estructura de las plantillas APX."""

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SHEETS = ROOT / "hojas-apx"


def dump(path: Path, max_row: int = 120, max_col: int = 14) -> None:
    print("=" * 100)
    print("FILE:", path.name)
    wb = load_workbook(path)
    print("SHEETS:", wb.sheetnames)
    for ws in wb.worksheets:
        print("-" * 100)
        print(f"SHEET '{ws.title}' dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        print("merged:", [str(m) for m in ws.merged_cells.ranges])
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_row), max_col=min(ws.max_column, max_col)):
            for cell in row:
                if cell.value is None:
                    continue
                val = str(cell.value).replace("\n", "\\n")
                if len(val) > 90:
                    val = val[:90] + "..."
                print(f"  {cell.coordinate:>6} | {val}")
        # anchos de columna
        widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
        print("  col widths:", widths)
        print("  data validations:")
        for dv in ws.data_validations.dataValidation:
            print("   ", dv.sqref, "->", dv.type, dv.formula1)


if __name__ == "__main__":
    files = sys.argv[1:] or [str(p) for p in SHEETS.glob("*.xlsx")]
    for f in files:
        dump(Path(f))
