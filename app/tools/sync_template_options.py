"""Sincroniza la hoja 'Parameters (Do not remove)' de las plantillas con opciones-*.json.

Es idempotente: se puede volver a ejecutar cuando cambien las listas del contrato.
"""

from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from apx_helper.contracts import LIBRARY, TRANSACTION, load_options
from apx_helper.settings import LIBRARY_TEMPLATE_GLOB, TRANSACTION_TEMPLATE_GLOB, find_template

PARAMS_SHEET = "Parameters (Do not remove)"

# Encabezado de la hoja Parameters -> clave de opciones-*.json que lo alimenta.
COLUMN_SOURCE = {
    TRANSACTION: {
        "Yes/No": "Asynchronous?",
        "Severity": "Severity",
        "APX Data Type": "APX data type",
        "Countries": "Country",
    },
    LIBRARY: {
        "Yes/No": "Mandatory?",
        "Visibility": "Visibility",
        "ONLINE_BATCH": "Library Type",
        "Access": "Access Type",
        "Format Type": "Format Type",
        "Countries": "Country",
    },
}


def sync(template: Path, kind: str, *, dry_run: bool = False) -> list[str]:
    options = load_options(kind)
    wb = load_workbook(template)
    ws = wb[PARAMS_SHEET]
    changes: list[str] = []

    headers = {
        str(ws.cell(row=1, column=col).value): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=col).value
    }

    for header, option_key in COLUMN_SOURCE[kind].items():
        col = headers.get(header)
        if col is None:
            changes.append(f"[!] La plantilla no tiene la columna '{header}'.")
            continue

        values = options[option_key]
        current = [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]
        current = [str(v) for v in current if v is not None]
        if current == values:
            continue

        style = copy(ws.cell(row=2, column=col)._style)
        for offset, value in enumerate(values):
            cell = ws.cell(row=2 + offset, column=col)
            cell.value = value
            cell._style = copy(style)
        for row in range(2 + len(values), ws.max_row + 1):
            ws.cell(row=row, column=col).value = None
        changes.append(f"    {header}: {current} -> {values}")

    if changes and not dry_run:
        wb.save(template)
    return changes


if __name__ == "__main__":
    dry_run = "--dry-run" in sys.argv
    for kind, pattern in ((TRANSACTION, TRANSACTION_TEMPLATE_GLOB), (LIBRARY, LIBRARY_TEMPLATE_GLOB)):
        template = find_template(pattern)
        changes = sync(template, kind, dry_run=dry_run)
        print(f"{template.name}: {'sin cambios' if not changes else ''}")
        for change in changes:
            print(change)
    if dry_run:
        print("\n(dry-run: no se escribió nada)")
