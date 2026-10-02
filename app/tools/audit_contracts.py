"""Audita la coherencia entre contratos JSON, mapeo de la app y desplegables de las plantillas."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from apx_helper import layout
from apx_helper.contracts import LIBRARY, TRANSACTION, load_options, load_schema
from apx_helper.settings import LIBRARY_TEMPLATE_GLOB, TRANSACTION_TEMPLATE_GLOB, find_template

PARAMS_SHEET = "Parameters (Do not remove)"


def schema_paths(node, prefix=()) -> set[tuple[str, ...]]:
    paths = set()
    if isinstance(node, dict):
        for key, value in node.items():
            paths |= schema_paths(value, prefix + (key,))
    elif isinstance(node, list):
        paths |= schema_paths(node[0] if node else {}, prefix + ("[]",))
    else:
        paths.add(prefix)
    return paths


def mapped_paths(fields, blocks) -> set[tuple[str, ...]]:
    paths = {f.path for f in fields}
    for block in blocks:
        for column in block.columns:
            if column.key:
                paths.add(block.path + ("[]", column.key))
    return paths


def template_lists(path: Path) -> dict[str, list[str]]:
    ws = load_workbook(path)[PARAMS_SHEET]
    columns = {}
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if not header:
            continue
        values = [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]
        columns[str(header)] = [str(v) for v in values if v is not None]
    return columns


def report(title: str, schema, fields, blocks, options, template, option_to_column, no_cell=frozenset()):
    print("=" * 90)
    print(title)
    missing = schema_paths(schema) - mapped_paths(fields, blocks) - no_cell
    extra = mapped_paths(fields, blocks) - schema_paths(schema)
    print("  Claves del contrato SIN celda asignada:", sorted(missing) or "ninguna")
    print("  Celdas que esperan claves INEXISTENTES:", sorted(extra) or "ninguna")

    used = {f.options for f in fields if f.options}
    for block in blocks:
        used |= {c.options for c in block.columns if c.options}
    print("  Listas usadas por la app sin definir en opciones-*.json:", sorted(used - set(options)) or "ninguna")
    print("  Listas definidas y no usadas:", sorted(set(options) - used) or "ninguna")

    lists = template_lists(template)
    for option_name, values in options.items():
        column = option_to_column.get(option_name)
        allowed = lists.get(column, [])
        unknown = [v for v in values if v not in allowed]
        if unknown:
            print(f"  [!] '{option_name}': {unknown} no existe(n) en el desplegable '{column}' -> {allowed}")


TRANSACTION_COLUMNS = {
    "Country": "Countries",
    "Asynchronous?": "Yes/No",
    "Transactional?": "Yes/No",
    "Severity": "Severity",
    "APX data type": "APX Data Type",
    "Mandatory?": "Yes/No",
}
LIBRARY_COLUMNS = {
    "Visibility": "Visibility",
    "Library Type": "ONLINE_BATCH",
    "Country": "Countries",
    "Access Type": "Access",
    "Format Type": "Format Type",
    "Mandatory?": "Yes/No",
}

if __name__ == "__main__":
    report(
        "TRANSACCIÓN",
        load_schema(TRANSACTION),
        layout.TRANSACTION_FIELDS,
        layout.TRANSACTION_BLOCKS,
        load_options(TRANSACTION),
        find_template(TRANSACTION_TEMPLATE_GLOB),
        TRANSACTION_COLUMNS,
    )
    library_fields = layout.LIBRARY_FIELDS + tuple(
        layout.Field(("EXECUTE - METHODS SUMMARY", "[]") + f.path, f.cell, f.options)
        for f in layout.METHOD_FIELDS
    )
    library_blocks = (
        layout.Block(
            ("EXECUTE - METHODS SUMMARY",),
            layout.LIBRARY_METHODS_BLOCK.first_row,
            layout.LIBRARY_METHODS_BLOCK.last_row,
            layout.LIBRARY_METHODS_BLOCK.columns,
        ),
    ) + tuple(
        layout.Block(
            ("EXECUTE - METHODS SUMMARY", "[]") + b.path, b.first_row, b.last_row, b.columns
        )
        for b in layout.METHOD_BLOCKS
    )
    report(
        "LIBRERÍA",
        load_schema(LIBRARY),
        library_fields,
        library_blocks,
        load_options(LIBRARY),
        find_template(LIBRARY_TEMPLATE_GLOB),
        LIBRARY_COLUMNS,
        no_cell={layout.LIBRARY_ID_PATH},
    )
