"""Rellena las plantillas Excel APX a partir de un contrato JSON validado."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from . import layout
from .contracts import LIBRARY, TRANSACTION, ContractError, load_options, validate
from .settings import LIBRARY_TEMPLATE_GLOB, TRANSACTION_TEMPLATE_GLOB, find_template, templates_dir
from .xlsx import (
    cell_ref,
    clear_cell,
    copy_sheet,
    insert_styled_rows,
    reorder_sheets,
    unique_sheet_title,
    write_cell,
)

REVIEW_MARKER = "[REVISAR"
MISSING_INFO = "NO HAY INFORMACIÓN SUFICIENTE"
_UNSAFE_FILENAME = re.compile(r"[^A-Za-z0-9._\- ]+")

MISSING = "missing"
REVIEW = "review"
INVALID = "invalid"

SEVERITY_LABELS = {
    MISSING: "Falta información",
    REVIEW: "A revisar",
    INVALID: "Valor no permitido",
}


@dataclass(frozen=True)
class Issue:
    """Un pendiente que el desarrollador debe completar o confirmar en la hoja."""

    sheet: str
    field: str
    severity: str
    message: str
    cell: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "cell": self.cell,
            "field": self.field,
            "severity": self.severity,
            "severityLabel": SEVERITY_LABELS[self.severity],
            "message": self.message,
        }

    def __str__(self) -> str:
        where = f"{self.sheet}!{self.cell}" if self.cell else self.sheet
        return f"{where} · {self.field}: {self.message}"


@dataclass
class RenderResult:
    workbook: Workbook
    filename: str
    issues: list[Issue] = field(default_factory=list)

    @property
    def warnings(self) -> list[str]:
        return [str(issue) for issue in self.issues]

    @property
    def summary(self) -> dict[str, int]:
        counts = {severity: 0 for severity in SEVERITY_LABELS}
        for issue in self.issues:
            counts[issue.severity] += 1
        return counts

    def save(self, directory: str | Path) -> Path:
        directory = Path(directory)
        directory.mkdir(parents=True, exist_ok=True)
        destination = (directory / self.filename).resolve()
        if destination.parent == templates_dir().resolve():
            raise ValueError("No se puede escribir la salida dentro de la carpeta de plantillas.")
        self.workbook.save(destination)
        return destination

    def to_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        return buffer.getvalue()


def safe_filename(name: str, default: str = "APX") -> str:
    clean = _UNSAFE_FILENAME.sub("_", str(name)).strip("._ ")
    return (clean or default)[:100]


def _output_filename(template: Path, identifier: str) -> str:
    """Convención usada por el equipo: '<identificador> - <nombre de la plantilla>.xlsx'."""
    return f"{safe_filename(identifier)} - {template.stem}.xlsx"


def _dig(data: Any, path: Sequence[str]) -> Any:
    node = data
    for key in path:
        if not isinstance(node, dict):
            return None
        node = node.get(key)
    return node


def _text(value: Any) -> str:
    if value is None or isinstance(value, str):
        return (value or "").strip()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)




class _Filler:
    def __init__(self, options: dict[str, list[str]]):
        self.options = options
        self.issues: list[Issue] = []

    def _add(self, ws: Worksheet, cell: str | None, field_label: str, severity: str, message: str) -> None:
        self.issues.append(Issue(ws.title, field_label, severity, message, cell))

    def _write(
        self,
        ws: Worksheet,
        row: int,
        col: int,
        coordinate: str,
        field_label: str,
        value: str,
        option_key: str | None,
    ) -> None:
        upper = value.upper()
        severity: str | None = None
        message = ""
        if MISSING_INFO in upper:
            severity, message = MISSING, "el analizador no encontró información suficiente en el código"
        elif REVIEW_MARKER in upper:
            severity, message = REVIEW, f"el analizador marcó el valor para confirmar: {value}"

        if option_key:
            allowed = self.options.get(option_key, [])
            if allowed and value not in allowed and severity is None:
                severity = INVALID
                message = f"'{value}' no está en la lista permitida ({', '.join(allowed)})"

        if severity:
            self._add(ws, coordinate, field_label, severity, message)
        write_cell(ws, row, col, value, highlight=severity is not None)

    def fields(self, ws: Worksheet, data: dict, fields: Iterable[layout.Field]) -> None:
        for spec in fields:
            field_label = " > ".join(spec.path)
            value = _text(_dig(data, spec.path))
            if not value:
                if not spec.optional:
                    self._add(ws, spec.cell, field_label, MISSING, "el contrato llega sin valor")
                continue
            row, col = cell_ref(spec.cell)
            self._write(ws, row, col, spec.cell, field_label, value, spec.options)

    def blocks(self, ws: Worksheet, data: dict, blocks: Iterable[layout.Block]) -> None:
        # De abajo hacia arriba: así una inserción no invalida las filas de los bloques
        # que aún no se han procesado.
        for block in sorted(blocks, key=lambda b: b.first_row, reverse=True):
            self.block(ws, _dig(data, block.path), block)

    def block(self, ws: Worksheet, raw_items: Any, block: layout.Block) -> None:
        items = _clean_items(raw_items, block)
        block_label = " > ".join(block.path)
        if not items and block.required:
            self._add(ws, None, block_label, MISSING, "el contrato no registra ningún elemento")

        rows = max(block.capacity, len(items))
        if len(items) > block.capacity:
            insert_styled_rows(ws, block.last_row + 1, len(items) - block.capacity, block.last_row)

        columns = [(column, column_index_from_string(column.letter)) for column in block.columns]
        for offset in range(rows):
            row = block.first_row + offset
            for _, col in columns:
                clear_cell(ws, row, col)

        for offset, item in enumerate(items):
            row = block.first_row + offset
            for column, col in columns:
                if column.literal is not None:
                    write_cell(ws, row, col, column.literal)
                    continue
                coordinate = f"{column.letter}{row}"
                field_label = f"{block_label} [{offset + 1}] > {column.key}"
                value = _text(item.get(column.key))
                if not value:
                    self._add(ws, coordinate, field_label, MISSING, "el contrato llega sin valor")
                    continue
                self._write(ws, row, col, coordinate, field_label, value, column.options)


def _clean_items(raw_items: Any, block: layout.Block) -> list[dict]:
    """Descarta los objetos plantilla vacíos que trae el contrato de referencia."""
    if not isinstance(raw_items, list):
        return []
    keys = [column.key for column in block.columns if column.key]
    items = []
    for item in raw_items:
        if not isinstance(item, dict):
            continue
        if any(_text(item.get(key)) for key in keys):
            items.append(item)
    return items


# --------------------------------------------------------------------------------------
# Transacción
# --------------------------------------------------------------------------------------


def render_transactions(contracts: Sequence[dict], *, name: str | None = None) -> RenderResult:
    """Genera un único libro con una hoja por transacción (una DU completa)."""
    if not contracts:
        raise ContractError("No se recibió ningún contrato de transacción.")

    template = find_template(TRANSACTION_TEMPLATE_GLOB)
    wb = load_workbook(template)
    base = wb.worksheets[0]
    # Las copias se hacen antes de escribir nada para que cada hoja parta de la plantilla limpia.
    sheets = [base] + [copy_sheet(wb, base, f"__apx_tmp_{i}") for i in range(1, len(contracts))]

    filler = _Filler(load_options(TRANSACTION))
    titles: list[str] = []
    for position, (ws, data) in enumerate(zip(sheets, contracts), start=1):
        identifier = _text(_dig(data, ("TECHNICAL DATA", "Transaction Identifier")))
        ws.title = unique_sheet_title(wb, identifier or f"TRANSACTION{position}", fallback=ws.title)
        titles.append(ws.title)
        filler.fields(ws, data, layout.TRANSACTION_FIELDS)
        filler.blocks(ws, data, layout.TRANSACTION_BLOCKS)

    reorder_sheets(wb, titles)
    return RenderResult(wb, _output_filename(template, name or titles[0]), filler.issues)


def render_transaction(data: dict, *, name: str | None = None) -> RenderResult:
    return render_transactions([data], name=name)


# --------------------------------------------------------------------------------------
# Librería
# --------------------------------------------------------------------------------------


def render_library(data: dict, *, library_id: str | None = None, name: str | None = None) -> RenderResult:
    template = find_template(LIBRARY_TEMPLATE_GLOB)
    wb = load_workbook(template)
    main = wb.worksheets[0]
    template_sheet = wb[layout.METHOD_TEMPLATE_SHEET]

    identifier = (
        library_id
        or name
        or _text(_dig(data, ("TECHNICAL DATA", "Library Identifier")))
        or main.title
    )
    main.title = unique_sheet_title(wb, identifier, fallback=main.title)

    filler = _Filler(load_options(LIBRARY))
    filler.fields(main, data, layout.LIBRARY_FIELDS)

    methods = _clean_items(data.get("EXECUTE - METHODS SUMMARY"), layout.LIBRARY_METHODS_BLOCK)
    filler.block(main, methods, layout.LIBRARY_METHODS_BLOCK)

    method_titles: list[str] = []
    for method in methods:
        title = _text(method.get("Method name")) or f"method{len(method_titles) + 1}"
        sheet = copy_sheet(wb, template_sheet, title)
        method_titles.append(sheet.title)
        filler.fields(sheet, method, layout.METHOD_FIELDS)
        filler.blocks(sheet, method, layout.METHOD_BLOCKS)

    if method_titles:
        wb.remove(template_sheet)

    reorder_sheets(wb, [main.title, *method_titles])
    return RenderResult(wb, _output_filename(template, identifier), filler.issues)


# --------------------------------------------------------------------------------------
# Entrada única
# --------------------------------------------------------------------------------------


def render(
    data: dict | Sequence[dict],
    *,
    kind: str | None = None,
    library_id: str | None = None,
    name: str | None = None,
) -> RenderResult:
    """Valida el/los contrato(s) y devuelve el Excel correspondiente ya relleno.

    Una lista de contratos de transacción se combina en un único libro (una hoja por
    transacción), tal y como se documenta una DU completa.
    """
    if isinstance(data, (list, tuple)):
        kinds = {validate(item) for item in data}
        if kinds - {TRANSACTION}:
            raise ContractError(
                "Solo se pueden combinar contratos de transacción en un mismo archivo; "
                "cada librería genera su propio libro."
            )
        return render_transactions(list(data), name=name)

    kind = validate(data, kind)
    if kind == TRANSACTION:
        return render_transaction(data, name=name)
    return render_library(data, library_id=library_id, name=name)
