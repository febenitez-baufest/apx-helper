"""Helpers de openpyxl para escribir en las plantillas sin alterar su formato."""

from __future__ import annotations

import re
from copy import copy
from typing import Any

from openpyxl.styles import PatternFill
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# Ámbar suave: marca celdas cuyo valor no cumple la lista desplegable o requiere revisión.
WARNING_FILL = PatternFill(fill_type="solid", start_color="FFFFE699", end_color="FFFFE699")

_INVALID_TITLE_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_SHEET_TITLE = 31


def sanitize_sheet_title(title: str, fallback: str = "Sheet") -> str:
    clean = _INVALID_TITLE_CHARS.sub("-", str(title)).strip().strip("'")
    clean = " ".join(clean.split())
    return clean[:MAX_SHEET_TITLE] or fallback


def unique_sheet_title(wb: Workbook, title: str, fallback: str = "Sheet") -> str:
    base = sanitize_sheet_title(title, fallback)
    candidate, index = base, 2
    while candidate in wb.sheetnames:
        suffix = f"_{index}"
        candidate = base[: MAX_SHEET_TITLE - len(suffix)] + suffix
        index += 1
    return candidate


def anchor_cell(ws: Worksheet, row: int, col: int) -> tuple[int, int]:
    """Devuelve la celda superior-izquierda del rango combinado que contiene (row, col)."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return rng.min_row, rng.min_col
    return row, col


def write_cell(ws: Worksheet, row: int, col: int, value: Any, *, highlight: bool = False):
    row, col = anchor_cell(ws, row, col)
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if highlight:
        cell.fill = copy(WARNING_FILL)
    return cell


def clear_cell(ws: Worksheet, row: int, col: int) -> None:
    row, col = anchor_cell(ws, row, col)
    ws.cell(row=row, column=col).value = None


def cell_ref(coordinate: str) -> tuple[int, int]:
    return coordinate_to_tuple(coordinate)


def insert_styled_rows(ws: Worksheet, index: int, amount: int, style_row: int) -> None:
    """Inserta `amount` filas antes de `index` clonando estilo, alto, combinaciones y
    validaciones de `style_row`. `style_row` debe estar por encima de `index`."""
    if amount <= 0:
        return
    if style_row >= index:
        raise ValueError("style_row debe estar por encima de la fila de inserción.")

    max_col = ws.max_column
    styles = [copy(ws.cell(row=style_row, column=c)._style) for c in range(1, max_col + 1)]
    height = ws.row_dimensions[style_row].height
    merge_cols = [
        (rng.min_col, rng.max_col)
        for rng in ws.merged_cells.ranges
        if rng.min_row == style_row and rng.max_row == style_row
    ]
    heights = {r: dim.height for r, dim in ws.row_dimensions.items() if dim.height is not None}

    ws.insert_rows(index, amount)

    for rng in ws.merged_cells.ranges:
        if rng.min_row >= index:
            rng.shift(row_shift=amount)
        elif rng.max_row >= index:
            rng.expand(down=amount)

    for validation in ws.data_validations.dataValidation:
        for rng in validation.sqref.ranges:
            if rng.min_row >= index:
                rng.shift(row_shift=amount)
            elif rng.max_row >= index or rng.max_row == style_row:
                rng.expand(down=amount)

    for row in sorted(heights, reverse=True):
        if row >= index:
            ws.row_dimensions[row + amount].height = heights[row]

    for offset in range(amount):
        row = index + offset
        if height is not None:
            ws.row_dimensions[row].height = height
        for col in range(1, max_col + 1):
            ws.cell(row=row, column=col)._style = copy(styles[col - 1])

    for min_col, max_col_ in merge_cols:
        for offset in range(amount):
            ws.merge_cells(
                start_row=index + offset,
                start_column=min_col,
                end_row=index + offset,
                end_column=max_col_,
            )


def copy_sheet(wb: Workbook, source: Worksheet, title: str) -> Worksheet:
    """Duplica una hoja incluyendo validaciones de datos y anchos de columna."""
    target = wb.copy_worksheet(source)
    target.title = unique_sheet_title(wb, title, fallback=source.title)

    for validation in source.data_validations.dataValidation:
        target.add_data_validation(copy(validation))
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dim)
        target.column_dimensions[key].worksheet = target
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key].height = dim.height
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes
    return target


def reorder_sheets(wb: Workbook, ordered_titles: list[str]) -> None:
    """Coloca las hojas indicadas al principio, conservando el resto en su orden."""
    front = [wb[title] for title in ordered_titles if title in wb.sheetnames]
    rest = [ws for ws in wb.worksheets if ws not in front]
    wb._sheets = front + rest
